# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

from io import BytesIO

import pytest
from openpyxl import load_workbook

from plane.utils.porters.formatters import XLSXFormatter


def _export_title_cell(name):
    """Export one row through XLSXFormatter and return its reopened Name cell"""
    content = XLSXFormatter().encode([{"name": name, "sequence_id": 1}])
    return load_workbook(BytesIO(content)).active["A2"]


@pytest.mark.unit
class TestXLSXFormatterFormulaInjection:
    """Test that exported values are never stored as live formulas"""

    def test_title_starting_with_equals_is_not_a_formula(self):
        """A title that looks like a formula is exported as text"""
        title = '=HYPERLINK("https://example.test","x")'
        cell = _export_title_cell(title)
        assert cell.data_type == "s"
        assert cell.value == title
        assert cell.quotePrefix is True

    @pytest.mark.parametrize("title", ["+1 for dark mode", "- follow up", "@mention bug", "\tindented"])
    def test_other_formula_triggers_are_quote_prefixed(self, title):
        """Titles Excel would evaluate on edit keep their text and are quote-prefixed"""
        cell = _export_title_cell(title)
        assert cell.data_type == "s"
        assert cell.value == title
        assert cell.quotePrefix is True

    def test_list_value_starting_with_equals_is_not_a_formula(self):
        """Joined list values (labels, assignees) are guarded too"""
        content = XLSXFormatter().encode([{"labels": ["=1+1", "bug"]}])
        cell = load_workbook(BytesIO(content)).active["A2"]
        assert cell.data_type == "s"
        assert cell.value == "=1+1, bug"

    def test_ordinary_title_is_unchanged(self):
        """Ordinary titles are not quote-prefixed"""
        cell = _export_title_cell("Fix login = broken")
        assert cell.data_type == "s"
        assert cell.value == "Fix login = broken"
        assert cell.quotePrefix is False

    def test_numbers_stay_numeric(self):
        """Non-string values keep their type"""
        content = XLSXFormatter().encode([{"name": "Task", "sequence_id": -5}])
        cell = load_workbook(BytesIO(content)).active["B2"]
        assert cell.data_type == "n"
        assert cell.value == -5
